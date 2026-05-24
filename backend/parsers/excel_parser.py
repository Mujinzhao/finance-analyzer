from __future__ import annotations

from pathlib import Path
from typing import Dict, Iterable

import openpyxl

BALANCE_SHEET_MAP: Dict[str, str] = {
    "货币资金": "cash_and_equivalents",
    "交易性金融资产": "trading_financial_assets",
    "应收票据": "notes_receivable",
    "应收账款": "accounts_receivable",
    "预付款项": "prepayments",
    "其他应收款": "other_receivables",
    "存货": "inventory",
    "其他流动资产": "other_current_assets",
    "合同资产": "contract_assets",
    "流动资产合计": "total_current_assets",
    "可供出售金融资产": "available_for_sale_assets",
    "其他权益工具投资": "available_for_sale_assets",
    "持有至到期投资": "held_to_maturity",
    "其他债权投资": "held_to_maturity",
    "投资性房地产": "investment_property",
    "使用权资产": "right_of_use_assets",
    "长期股权投资": "long_term_equity_investment",
    "固定资产": "fixed_assets",
    "在建工程": "construction_in_progress",
    "无形资产": "intangible_assets",
    "商誉": "goodwill",
    "长期待摊费用": "long_term_deferred_expenses",
    "递延所得税资产": "deferred_tax_assets",
    "其他非流动资产": "other_non_current_assets",
    "非流动资产合计": "total_non_current_assets",
    "资产总计": "total_assets",
    "短期借款": "short_term_borrowings",
    "应付票据": "notes_payable",
    "应付账款": "accounts_payable",
    "预收款项": "advance_receipts",
    "合同负债": "advance_receipts",
    "应付职工薪酬": "employee_compensation_payable",
    "应交税费": "taxes_payable",
    "其他流动负债": "other_current_liabilities",
    "一年内到期的非流动负债": "non_current_liabilities_due_within_one_year",
    "其他应付款": "other_payables",
    "流动负债合计": "total_current_liabilities",
    "长期借款": "long_term_borrowings",
    "应付债券": "bonds_payable",
    "递延所得税负债": "deferred_tax_liabilities",
    "其他非流动负债": "other_non_current_liabilities",
    "长期应付款": "long_term_payables",
    "租赁负债": "lease_liabilities",
    "递延收益": "deferred_income",
    "非流动负债合计": "total_non_current_liabilities",
    "负债合计": "total_liabilities",
    "实收资本": "paid_in_capital",
    "股本": "paid_in_capital",
    "资本公积": "capital_reserve",
    "盈余公积": "surplus_reserve",
    "未分配利润": "undistributed_profit",
    "归属于母公司所有者权益合计": "total_equity_parent",
    "归属于母公司股东权益合计": "total_equity_parent",
    "其他综合收益": "other_comprehensive_income",
    "少数股东权益": "minority_interest",
    "所有者权益合计": "total_equity",
    "股东权益合计": "total_equity",
}

INCOME_STATEMENT_MAP: Dict[str, str] = {
    "营业收入": "revenue",
    "营业成本": "cost_of_revenue",
    "税金及附加": "taxes_and_surcharges",
    "销售费用": "selling_expenses",
    "管理费用": "admin_expenses",
    "研发费用": "rd_expenses",
    "财务费用": "finance_expenses",
    "利息费用": "interest_expense",
    "利息收入": "interest_income",
    "资产减值损失": "asset_impairment_loss",
    "信用减值损失": "credit_impairment_loss",
    "其他收益": "other_income",
    "投资收益": "investment_income",
    "公允价值变动收益": "fair_value_change_income",
    "资产处置收益": "asset_disposal_income",
    "营业利润": "operating_profit",
    "营业外收入": "non_operating_income",
    "营业外支出": "non_operating_expenses",
    "利润总额": "total_profit",
    "所得税费用": "income_tax_expense",
    "净利润": "net_profit",
    "归属于母公司股东的净利润": "net_profit_parent",
    "归属于母公司所有者的净利润": "net_profit_parent",
    "扣除非经常性损益后的净利润": "non_recurring_profit",
    "扣除非经常性损益的净利润": "non_recurring_profit",
    "扣非净利润": "non_recurring_profit",
    "少数股东损益": "net_profit_minority",
    "总股本(万股)": "total_shares",
    "总股本": "total_shares",
    "分红总额": "dividends_paid",
}

CASH_FLOW_MAP: Dict[str, str] = {
    "销售商品、提供劳务收到的现金": "cash_from_sales",
    "收到的税费返还": "tax_refunds",
    "收到其他与经营活动有关的现金": "other_cash_from_operations",
    "经营活动现金流入小计": "total_cash_from_operations",
    "购买商品、接受劳务支付的现金": "cash_paid_for_goods",
    "支付给职工以及为职工支付的现金": "cash_paid_to_employees",
    "支付的各项税费": "taxes_paid",
    "支付其他与经营活动有关的现金": "other_cash_paid_for_operations",
    "经营活动现金流出小计": "total_cash_paid_for_operations",
    "经营活动产生的现金流量净额": "net_cash_from_operations",
    "收回投资收到的现金": "cash_from_investment_withdrawal",
    "取得投资收益收到的现金": "cash_from_investment_income",
    "处置固定资产等收到的现金净额": "cash_from_asset_disposal",
    "处置固定资产、无形资产和其他长期资产收回的现金净额": "cash_from_asset_disposal",
    "收到其他与投资活动有关的现金": "other_cash_from_investing",
    "投资活动现金流入小计": "total_cash_from_investing",
    "购建固定资产、无形资产和其他长期资产支付的现金": "cash_paid_for_assets",
    "投资支付的现金": "cash_paid_for_investments",
    "支付其他与投资活动有关的现金": "other_cash_paid_for_investing",
    "投资活动现金流出小计": "total_cash_paid_for_investing",
    "投资活动产生的现金流量净额": "net_cash_from_investing",
    "取得借款收到的现金": "cash_from_borrowings",
    "吸收投资收到的现金": "cash_from_equity_issuance",
    "收到其他与筹资活动有关的现金": "other_cash_from_financing",
    "筹资活动现金流入小计": "total_cash_from_financing",
    "偿还债务支付的现金": "cash_paid_for_debt",
    "分配股利、利润或偿付利息支付的现金": "cash_paid_for_dividends",
    "支付其他与筹资活动有关的现金": "other_cash_paid_for_financing",
    "筹资活动现金流出小计": "total_cash_paid_for_financing",
    "筹资活动产生的现金流量净额": "net_cash_from_financing",
    "现金及现金等价物净增加额": "net_increase_in_cash",
    "期初现金及现金等价物余额": "cash_at_beginning",
    "期末现金及现金等价物余额": "cash_at_end",
}

ALL_MAPS: Iterable[Dict[str, str]] = (BALANCE_SHEET_MAP, INCOME_STATEMENT_MAP, CASH_FLOW_MAP)


def _to_number(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return 0.0
    negative = text.startswith("(") and text.endswith(")")
    text = text.replace(",", "").replace("，", "")
    text = text.replace("(", "").replace(")", "")
    multiplier = 10000.0 if "万" in text and "万股" not in text else 1.0
    text = text.replace("万元", "").replace("万", "").replace("元", "")
    try:
        number = float(text) * multiplier
    except ValueError:
        return 0.0
    return -number if negative else number


def parse_excel(file_path: str | Path, company_id: int, year: int, quarter: int = 0) -> Dict[str, float]:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    data: Dict[str, float] = {}

    for sheet in wb.worksheets:
        for row in sheet.iter_rows(min_row=1, max_col=2, values_only=True):
            name = str(row[0]).strip() if row[0] is not None else ""
            amount = row[1]
            if not name:
                continue
            for mapping in ALL_MAPS:
                field = mapping.get(name)
                if field and field not in data:
                    data[field] = _to_number(amount)

    if "基本信息" in wb.sheetnames:
        info_sheet = wb["基本信息"]
        for row in info_sheet.iter_rows(min_row=1, max_col=2, values_only=True):
            key = str(row[0]).strip() if row[0] is not None else ""
            value = row[1]
            if key in ("总股本", "总股本(万股)") and "total_shares" not in data:
                data["total_shares"] = _to_number(value)
            if key == "分红总额" and "dividends_paid" not in data:
                data["dividends_paid"] = _to_number(value)

    data["company_id"] = company_id
    data["year"] = year
    data["quarter"] = quarter
    data["report_type"] = "annual" if quarter == 0 else "quarterly"
    return data


def generate_template(path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    basic = wb.create_sheet("基本信息")
    basic.append(["科目", "金额"])
    basic.append(["总股本(万股)", ""])
    basic.append(["分红总额", ""])

    mapping_sheets = [
        ("资产负债表", BALANCE_SHEET_MAP),
        ("利润表", INCOME_STATEMENT_MAP),
        ("现金流量表", CASH_FLOW_MAP),
    ]

    for sheet_name, mapping in mapping_sheets:
        ws = wb.create_sheet(sheet_name)
        ws.append(["科目", "金额"])
        seen = set()
        for key in mapping.keys():
            if key in seen:
                continue
            seen.add(key)
            ws.append([key, ""])

    wb.save(path)
    return path
